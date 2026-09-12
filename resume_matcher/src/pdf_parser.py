"""
pdf_parser.py
-------------
Extracts raw text from resume/JD files. Despite the module name (kept for
import-compatibility with the rest of the pipeline), this now handles three
formats:
  - .pdf  : pdfplumber, falling back to pypdf if the first pass fails/returns
            too little text
  - .docx : python-docx, reading paragraphs and table cells
  - .txt  : read directly (useful for testing without generating real files)

Strategy for PDFs:
1. Try pdfplumber (best layout/text fidelity).
2. Fall back to pypdf if pdfplumber fails or returns empty text.
3. Raise a clear error if the PDF is scanned/image-only (no extractable text),
   since OCR is out of scope for this module.

Any other extension found in a directory (.doc, .xml, .rtf, etc.) is skipped
LOUDLY - printed as a warning, never silently dropped - since a resume that
silently vanishes from the ranking is worse than one that visibly failed.
"""

from __future__ import annotations
import os
from dataclasses import dataclass

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt")


class PDFParsingError(Exception):
    """Raised when a file cannot be parsed into usable text."""


@dataclass
class ParsedPDF:
    path: str
    filename: str
    text: str
    num_pages: int


def _extract_with_pdfplumber(path: str) -> tuple[str, int]:
    import pdfplumber

    text_parts = []
    with pdfplumber.open(path) as pdf:
        num_pages = len(pdf.pages)
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text_parts.append(page_text)
    return "\n".join(text_parts), num_pages


def _extract_with_pypdf(path: str) -> tuple[str, int]:
    from pypdf import PdfReader

    reader = PdfReader(path)
    num_pages = len(reader.pages)
    text_parts = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(text_parts), num_pages


def _extract_with_docx(path: str) -> tuple[str, int]:
    """
    Extract text from a .docx file: paragraphs, plus any text sitting in
    tables (resumes built from templates often put skills/dates in table
    cells rather than plain paragraphs).
    """
    import docx

    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    text = "\n".join(parts)
    # .docx has no fixed "page" concept; report 1 as a placeholder so the
    # ParsedPDF shape stays uniform across formats.
    return text, 1


def parse_pdf(path: str, min_chars: int = 20) -> ParsedPDF:
    """
    Extract text from a single PDF file.

    Raises PDFParsingError if no usable text is found (e.g. a scanned
    image-only PDF that would need OCR).
    """
    if not os.path.isfile(path):
        raise PDFParsingError(f"File not found: {path}")

    filename = os.path.basename(path)
    text, num_pages = "", 0

    try:
        text, num_pages = _extract_with_pdfplumber(path)
    except Exception:
        pass

    if len(text.strip()) < min_chars:
        try:
            text2, num_pages2 = _extract_with_pypdf(path)
            if len(text2.strip()) > len(text.strip()):
                text, num_pages = text2, num_pages2
        except Exception:
            pass

    if len(text.strip()) < min_chars:
        raise PDFParsingError(
            f"Could not extract text from '{filename}'. It may be a scanned "
            f"image-only PDF that requires OCR, which this pipeline does not perform."
        )

    return ParsedPDF(path=path, filename=filename, text=text, num_pages=num_pages)


def parse_docx(path: str, min_chars: int = 20) -> ParsedPDF:
    """Extract text from a single .docx file."""
    if not os.path.isfile(path):
        raise PDFParsingError(f"File not found: {path}")

    filename = os.path.basename(path)
    try:
        text, num_pages = _extract_with_docx(path)
    except Exception as e:
        raise PDFParsingError(f"Could not read '{filename}' as a .docx file: {e}")

    if len(text.strip()) < min_chars:
        raise PDFParsingError(
            f"'{filename}' produced almost no text ({len(text.strip())} chars) - "
            f"it may be empty, corrupted, or an old binary .doc file saved with "
            f"a .docx extension."
        )

    return ParsedPDF(path=path, filename=filename, text=text, num_pages=num_pages)


def parse_pdf_or_text(path: str) -> ParsedPDF:
    """
    Dispatch by extension: .pdf -> parse_pdf, .docx -> parse_docx,
    anything else -> read as plain text (useful for .txt test fixtures).
    """
    lower = path.lower()
    if lower.endswith(".pdf"):
        return parse_pdf(path)
    if lower.endswith(".docx"):
        return parse_docx(path)
    with open(path, "r", encoding="utf-8") as f:
        text = f.read()
    return ParsedPDF(path=path, filename=os.path.basename(path), text=text, num_pages=1)


def parse_directory(dir_path: str) -> list[ParsedPDF]:
    """
    Parse every .pdf/.docx/.txt file in a directory (non-recursive), sorted
    by filename. Any other extension is skipped but reported as a warning -
    never silently dropped - so a candidate never vanishes from the ranking
    without a visible trace.
    """
    results = []
    errors = []
    skipped = []

    for fname in sorted(os.listdir(dir_path)):
        ext = os.path.splitext(fname)[1].lower()
        full_path = os.path.join(dir_path, fname)
        if not os.path.isfile(full_path):
            continue
        if ext not in SUPPORTED_EXTENSIONS:
            skipped.append(fname)
            continue
        try:
            results.append(parse_pdf_or_text(full_path))
        except PDFParsingError as e:
            errors.append(str(e))

    if skipped:
        print(f"[pdf_parser] Skipped {len(skipped)} unsupported file(s) "
              f"(supported: {', '.join(SUPPORTED_EXTENSIONS)}):")
        for fname in skipped:
            print(f"  - {fname}")

    if errors:
        print("[pdf_parser] Warnings while parsing directory:")
        for e in errors:
            print(f"  - {e}")

    return results